#!/usr/bin/env python3
"""
compare_planners_interactive.py

Compare Nav2 SmacPlannerHybrid (via /compute_path_to_pose) against your
a_star_node.py (via /goal_pose -> /path) on start/goal pairs YOU pick in RViz.

RViz workflow
  1. "2D Pose Estimate"  -> sets the start (published to /initialpose)
  2. "2D Nav Goal"       -> sets the goal  (published to /goal_pose)
     That one click triggers your a_star_node exactly as in normal use; this
     script sees the same message and then queries Nav2 with the identical pair.
  Repeat step 2 for new goals from the same start, or step 1 to move the start.
  Add ONE "MarkerArray" display on /compare_markers (orange = Nav2, blue = yours).

Everything is appended to ~/planner_comparisons.xlsx, three sheets:
  trials  one row per trial: start/goal, metrics for both planners, config_id
  paths   every waypoint of both output paths (trial, planner, idx, x, y, yaw_deg)
  configs live parameter dump of planner_server, global_costmap and a_star,
          written once per distinct configuration (keyed by config_id)
Trial numbering continues from an existing workbook, so the log is continuous.

Terminals
  1: your a_star_node
  2: ros2 launch compare_bringup.launch.py map_yaml:=/abs/path/map.yaml
  3: python3 compare_planners_interactive.py

Design notes (why it differs from the previous version)
  * Planner work runs in a worker thread; the ROS executor is multi-threaded.
    (Calling spin_until_future_complete inside a subscription callback of a
    node that is already spinning fails/deadlocks.)
  * Yours is collected FIRST, then Nav2 runs -- otherwise both planners
    compute at the same time on one machine and skew each other's timings.
  * Your planner's time = search (+ heuristic map) + smoothing, parsed from
    /rosout, because Nav2's planning_time already includes its smoother.
  * Parameters are read live over the ROS parameter services, so what is in the
    sheet is what actually ran, not what the yaml says.
"""
import hashlib
import json
import math
import os
import queue
import re
import threading
import time
import traceback
from datetime import datetime

import rclpy
from rclpy.node import Node
from rclpy.action import ActionClient
from rclpy.callback_groups import ReentrantCallbackGroup
from rclpy.duration import Duration
from rclpy.executors import MultiThreadedExecutor
from rclpy.qos import QoSProfile, DurabilityPolicy, ReliabilityPolicy, HistoryPolicy

from geometry_msgs.msg import PoseStamped, PoseWithCovarianceStamped, TransformStamped, Point, Quaternion
from nav_msgs.msg import Path
from std_msgs.msg import ColorRGBA
from visualization_msgs.msg import Marker, MarkerArray
from rcl_interfaces.msg import Log
from rcl_interfaces.srv import ListParameters, GetParameters
from tf2_ros import TransformBroadcaster
from tf_transformations import euler_from_quaternion
from nav2_msgs.action import ComputePathToPose

from openpyxl import Workbook, load_workbook

# Optional batch of pairs run automatically before interactive mode.
# CUSTOM_RUNS = {
#     "my_pair_1": {"start": (-8.2, -8.2, math.radians(45)), "goal": (8.5, 8.5, math.radians(45))},
# }
CUSTOM_RUNS = {}

NAV2_ERROR_CODES = {
    0: "NONE", 200: "UNKNOWN", 201: "INVALID_PLANNER", 202: "TF_ERROR",
    203: "START_OUTSIDE_MAP", 204: "GOAL_OUTSIDE_MAP", 205: "START_OCCUPIED",
    206: "GOAL_OCCUPIED", 207: "TIMEOUT", 208: "NO_VALID_PATH",
}

GOAL_TOPIC = "/goal_pose"
INITIALPOSE_TOPIC = "/initialpose"
YOURS_PATH_TOPIC = "/path"          # a_star_node's output -- subscribe only
MARKER_TOPIC = "/compare_markers"
ROBOT_FRAME = "zippy"
MAP_FRAME = "map"
YOUR_NODE_NAME = "a_star"           # logger name in /rosout
YOUR_PLANNER_TIMEOUT_S = 60.0
NAV2_ACTION_TIMEOUT_S = 40.0
START_SETTLE_S = 1.0

# ROS nodes whose live parameters are recorded per configuration.
PARAM_TARGETS = {
    "nav2_planner": "/planner_server",
    "nav2_costmap": "/global_costmap/global_costmap",
    "yours": "/a_star",
}

EXCEL_PATH = os.path.expanduser("~/planner_comparisons.xlsx")
SHEET_TRIALS, SHEET_PATHS, SHEET_CONFIGS = "trials", "paths", "configs"
HEADER = [
    "trial", "label", "timestamp", "config_id",
    "start_x", "start_y", "start_yaw_deg", "goal_x", "goal_y", "goal_yaw_deg", "straight_line_m",
    "nav2_success", "nav2_path_length_m", "nav2_num_waypoints",
    "nav2_planning_time_s", "nav2_wall_time_s", "nav2_error",
    "yours_success", "yours_path_length_m", "yours_num_waypoints",
    "yours_search_time_s", "yours_smoothing_time_s", "yours_total_time_s", "yours_wall_time_s",
    "yours_time_is_wallclock_fallback", "yours_expansions", "yours_error",
    "length_diff_m_yours_minus_nav2", "planning_time_diff_s_yours_minus_nav2",
]
PATH_HEADER = ["trial", "planner", "idx", "x", "y", "yaw_deg"]
CONFIG_HEADER = ["config_id", "first_seen", "source", "param", "value"]


# ----------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------
def yaw_to_quat(yaw):
    return (0.0, 0.0, math.sin(yaw / 2.0), math.cos(yaw / 2.0))


def quat_to_yaw(q):
    return euler_from_quaternion([q.x, q.y, q.z, q.w])[2]


def path_length(poses):
    return sum(math.hypot(b.pose.position.x - a.pose.position.x,
                          b.pose.position.y - a.pose.position.y)
               for a, b in zip(poses, poses[1:]))


def rgba(r, g, b, a=1.0):
    c = ColorRGBA()
    c.r, c.g, c.b, c.a = r, g, b, a
    return c


def rnd(v, nd):
    return round(v, nd) if isinstance(v, (int, float)) and not isinstance(v, bool) else ("" if v is None else v)


def fmt(v, nd=3):
    return f"{v:.{nd}f}" if isinstance(v, (int, float)) else "n/a"


def param_value(v):
    """rcl_interfaces/ParameterValue -> python value."""
    return {1: v.bool_value, 2: v.integer_value, 3: v.double_value, 4: v.string_value,
            5: list(v.byte_array_value), 6: list(v.bool_array_value),
            7: list(v.integer_array_value), 8: list(v.double_array_value),
            9: list(v.string_array_value)}.get(v.type)


class CompareNode(Node):
    def __init__(self):
        super().__init__("planner_compare_interactive")
        self.cb = ReentrantCallbackGroup()

        self.tf_broadcaster = TransformBroadcaster(self)
        self.goal_pub = self.create_publisher(PoseStamped, GOAL_TOPIC, 10)
        self.marker_pub = self.create_publisher(MarkerArray, MARKER_TOPIC, 10)

        path_qos = QoSProfile(depth=1)
        path_qos.durability = DurabilityPolicy.TRANSIENT_LOCAL
        path_qos.reliability = ReliabilityPolicy.RELIABLE
        path_qos.history = HistoryPolicy.KEEP_LAST
        self._last_path_msg = None
        self.create_subscription(Path, YOURS_PATH_TOPIC, self._path_cb, path_qos,
                                 callback_group=self.cb)

        self._alog = {}
        self._reset_yours_state()
        self.create_subscription(Log, "/rosout", self._rosout_cb, 100, callback_group=self.cb)

        self.nav2_client = ActionClient(self, ComputePathToPose, "compute_path_to_pose",
                                        callback_group=self.cb)
        self._param_clients = {}

        self.pending_start = None
        self.tf_x = 0.0
        self.tf_y = 0.0
        self.tf_yaw = 0.0
        self._start_set_time = 0.0
        self._batch_publishing = False
        self._goal_t0 = 0.0
        self._last_yours_goal_ok = None
        self.trial_counter = 0
        self._known_configs = set()

        self._jobs = queue.Queue()
        self._busy = threading.Event()
        self._stop = threading.Event()
        self._worker_thread = threading.Thread(target=self._worker, daemon=True)

        self.create_subscription(PoseWithCovarianceStamped, INITIALPOSE_TOPIC,
                                 self._initialpose_cb, 10, callback_group=self.cb)
        self.create_subscription(PoseStamped, GOAL_TOPIC, self._goal_cb, 10,
                                 callback_group=self.cb)

        # Nav2's global costmap waits for map->zippy TF before it activates, so
        # give it one right away (replaced as soon as you click a start pose).
        self.broadcast_start_tf(0.0, 0.0, 0.0)
        self.tf_timer = self.create_timer(0.05, self.publish_tf)
        self._clear_markers()
        self._ensure_workbook()
        self.get_logger().info(
            "Ready. RViz: '2D Pose Estimate' (start) then '2D Nav Goal' (goal). "
            f"Logging to {EXCEL_PATH} (continuing at trial {self.trial_counter + 1}).")

    def start(self):
        self._worker_thread.start()

    def stop(self):
        self._stop.set()

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------
    def _new_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_TRIALS
        ws.append(HEADER)
        wb.create_sheet(SHEET_PATHS).append(PATH_HEADER)
        wb.create_sheet(SHEET_CONFIGS).append(CONFIG_HEADER)
        wb.save(EXCEL_PATH)

    def _ensure_workbook(self):
        if os.path.exists(EXCEL_PATH):
            try:
                wb = load_workbook(EXCEL_PATH)
                ok = (all(s in wb.sheetnames for s in (SHEET_TRIALS, SHEET_PATHS, SHEET_CONFIGS))
                      and [c.value for c in wb[SHEET_TRIALS][1]] == HEADER)
                if ok:
                    self._known_configs = {r[0] for r in wb[SHEET_CONFIGS].iter_rows(min_row=2, values_only=True) if r[0]}
                    nums = [r[0] for r in wb[SHEET_TRIALS].iter_rows(min_row=2, values_only=True)
                            if isinstance(r[0], int)]
                    self.trial_counter = max(nums) if nums else 0
                    return
            except Exception as e:
                self.get_logger().warn(f"Could not read existing workbook: {e}")
            backup = EXCEL_PATH.replace(".xlsx", f".old_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
            os.rename(EXCEL_PATH, backup)
            self.get_logger().warn(f"Existing workbook had a different layout; moved to {backup}")
        self._new_workbook()

    def _save(self, trial_row, path_rows, config_id, snap):
        try:
            wb = load_workbook(EXCEL_PATH)
            wb[SHEET_TRIALS].append([trial_row.get(c, "") for c in HEADER])
            for r in path_rows:
                wb[SHEET_PATHS].append(r)
            new_cfg = config_id not in self._known_configs
            if new_cfg:
                ts = datetime.now().isoformat(timespec="seconds")
                for src, params in snap.items():
                    for k, v in sorted(params.items()):
                        wb[SHEET_CONFIGS].append([config_id, ts, src, k,
                                                  json.dumps(v) if isinstance(v, (list, dict)) else v])
            wb.save(EXCEL_PATH)
            if new_cfg:
                self._known_configs.add(config_id)
            return True
        except Exception as e:
            backup = os.path.expanduser(f"~/planner_trial_{trial_row.get('trial')}_backup.json")
            with open(backup, "w") as f:
                json.dump({"trial": trial_row, "paths": path_rows, "config_id": config_id, "config": snap},
                          f, indent=1, default=str)
            self.get_logger().error(
                f"Excel write failed ({e}) -- is the file open in Excel/LibreOffice? "
                f"Trial saved to {backup} instead.")
            return False

    # ------------------------------------------------------------------
    # Live parameter snapshot
    # ------------------------------------------------------------------
    def _client(self, srv_type, target, name):
        key = (target, name)
        if key not in self._param_clients:
            self._param_clients[key] = self.create_client(srv_type, f"{target}/{name}",
                                                          callback_group=self.cb)
        return self._param_clients[key]

    def _call(self, client, req, timeout):
        fut = client.call_async(req)
        end = time.perf_counter() + timeout
        while not fut.done() and time.perf_counter() < end:
            time.sleep(0.005)
        return fut.result() if fut.done() else None

    def fetch_params(self, target, timeout=3.0):
        lc = self._client(ListParameters, target, "list_parameters")
        gc = self._client(GetParameters, target, "get_parameters")
        if not lc.wait_for_service(timeout_sec=timeout) or not gc.wait_for_service(timeout_sec=timeout):
            return None
        req = ListParameters.Request()
        req.depth = ListParameters.Request.DEPTH_RECURSIVE
        lres = self._call(lc, req, timeout)
        if lres is None:
            return None
        names = [n for n in lres.result.names if not n.startswith("qos_overrides")]
        greq = GetParameters.Request()
        greq.names = names
        gres = self._call(gc, greq, timeout)
        if gres is None:
            return None
        return {n: param_value(v) for n, v in zip(names, gres.values)}

    def snapshot_config(self):
        snap = {}
        for src, target in PARAM_TARGETS.items():
            p = self.fetch_params(target)
            if p is None:
                self.get_logger().warn(f"Could not read parameters of {target}")
                p = {"_unavailable": True}
            snap[src] = p
        cid = hashlib.sha1(json.dumps(snap, sort_keys=True, default=str).encode()).hexdigest()[:8]
        return cid, snap

    # ------------------------------------------------------------------
    # Subscribers
    # ------------------------------------------------------------------
    def _reset_yours_state(self):
        self._last_path_msg = None
        self._alog = {"plan": None, "smooth": None, "expansions": None, "nopath": False}

    def _path_cb(self, msg):
        self._last_path_msg = msg

    def _rosout_cb(self, msg):
        if msg.name.split(".")[-1] != YOUR_NODE_NAME:
            return
        t = msg.msg
        try:
            if t.startswith("Path planning time"):
                self._alog["plan"] = float(t.split(":")[1].strip().split()[0])
            elif t.startswith("Smoothing:"):
                m = re.search(r"in ([0-9.eE+-]+) s", t)
                if m:
                    self._alog["smooth"] = float(m.group(1))
            elif t.startswith("Expansions:"):
                m = re.match(r"Expansions: (\d+)", t)
                if m:
                    self._alog["expansions"] = int(m.group(1))
            elif "No path found" in t:
                self._alog["nopath"] = True
        except (IndexError, ValueError):
            pass

    def _initialpose_cb(self, msg):
        if self._busy.is_set():
            self.get_logger().warn("Trial in progress -- start pose ignored.")
            return
        x, y = msg.pose.pose.position.x, msg.pose.pose.position.y
        yaw = quat_to_yaw(msg.pose.pose.orientation)
        self.pending_start = (x, y, yaw)
        self._start_set_time = time.perf_counter()
        self.broadcast_start_tf(x, y, yaw)
        self._publish_start_marker(x, y, yaw)
        self.get_logger().info(f"Start set: ({x:.2f}, {y:.2f}, {math.degrees(yaw):.0f} deg). Now click '2D Nav Goal'.")

    def _goal_cb(self, msg):
        if self._batch_publishing:          # our own batch-mode publish
            self._batch_publishing = False
            return
        if self._busy.is_set():
            self.get_logger().warn("Trial in progress -- goal ignored.")
            return
        if self.pending_start is None:
            self.get_logger().warn("No start set yet -- click '2D Pose Estimate' first.")
            return
        now = time.perf_counter()
        if now - self._start_set_time < START_SETTLE_S:
            self.get_logger().warn(
                "Goal clicked <1 s after the start; a_star_node may have planned from the previous start.")
        self._busy.set()
        self._reset_yours_state()           # BEFORE anything else, so its path can't be wiped later
        self._goal_t0 = now
        goal = (msg.pose.position.x, msg.pose.position.y, quat_to_yaw(msg.pose.orientation))
        self._publish_goal_marker(*goal)
        self._jobs.put((f"interactive_{self.trial_counter + 1}", self.pending_start, goal, False))

    # ------------------------------------------------------------------
    # TF / pose helpers
    # ------------------------------------------------------------------
    def broadcast_start_tf(self, x, y, yaw):
        self.tf_x = x
        self.tf_y = y
        self.tf_yaw = yaw
        
    def publish_tf(self):
        t = TransformStamped()

        t.header.stamp = self.get_clock().now().to_msg()
        t.header.frame_id = MAP_FRAME
        t.child_frame_id = ROBOT_FRAME

        t.transform.translation.x = self.tf_x
        t.transform.translation.y = self.tf_y
        t.transform.translation.z = 0.0

        qx, qy, qz, qw = yaw_to_quat(self.tf_yaw)

        t.transform.rotation.x = qx
        t.transform.rotation.y = qy
        t.transform.rotation.z = qz
        t.transform.rotation.w = qw

        self.tf_broadcaster.sendTransform(t)
    def make_pose(self, x, y, yaw):
        p = PoseStamped()
        p.header.frame_id = MAP_FRAME
        p.header.stamp = self.get_clock().now().to_msg()
        p.pose.position.x, p.pose.position.y = x, y
        qx, qy, qz, qw = yaw_to_quat(yaw)
        p.pose.orientation.x, p.pose.orientation.y = qx, qy
        p.pose.orientation.z, p.pose.orientation.w = qz, qw
        return p

    # ------------------------------------------------------------------
    # Planner calls (worker thread)
    # ------------------------------------------------------------------
    def _wait(self, fut, timeout):
        end = time.perf_counter() + timeout
        while not fut.done() and time.perf_counter() < end:
            time.sleep(0.005)
        return fut.done()

    def run_nav2(self, start, goal):
        if not self.nav2_client.wait_for_server(timeout_sec=10.0):
            return {"success": False, "error": "compute_path_to_pose server not available "
                                               "(is the costmap active? needs map->zippy TF)"}
        gm = ComputePathToPose.Goal()
        gm.goal = self.make_pose(*goal)
        gm.start = self.make_pose(*start)
        gm.planner_id = "GridBased"
        try:
            gm.use_start = True
        except AttributeError:
            pass

        t0 = time.perf_counter()
        sf = self.nav2_client.send_goal_async(gm)
        if not self._wait(sf, NAV2_ACTION_TIMEOUT_S) or not sf.result().accepted:
            return {"success": False, "error": "goal rejected or action server timeout"}
        rf = sf.result().get_result_async()
        if not self._wait(rf, NAV2_ACTION_TIMEOUT_S):
            return {"success": False, "error": "result timeout"}
        wall = time.perf_counter() - t0

        result = rf.result().result
        code = getattr(result, "error_code", 0)
        poses = list(result.path.poses)
        ok = len(poses) > 0 and code == 0
        pt = None
        if hasattr(result, "planning_time"):
            pt = result.planning_time.sec + result.planning_time.nanosec * 1e-9
        return {"success": ok, "path_length_m": path_length(poses) if ok else None,
                "num_waypoints": len(poses), "planning_time_s": pt, "wall_time_s": wall,
                "error": "" if ok else NAV2_ERROR_CODES.get(code, str(code)),
                "poses": poses if ok else []}

    def collect_yours(self):
        deadline = self._goal_t0 + YOUR_PLANNER_TIMEOUT_S
        while time.perf_counter() < deadline:
            if self._last_path_msg is not None:
                break
            if self._alog["nopath"]:
                return {"success": False, "error": "a_star_node logged 'No path found!'",
                        "wall_time_s": time.perf_counter() - self._goal_t0,
                        "search_time_s": self._alog["plan"], "expansions": self._alog["expansions"]}
            time.sleep(0.01)
        else:
            return {"success": False, "error": f"no /path within {YOUR_PLANNER_TIMEOUT_S:.0f}s"}
        wall = time.perf_counter() - self._goal_t0

        # /rosout is asynchronous: the log lines can trail the /path message.
        end = time.perf_counter() + 1.5
        while self._alog["plan"] is None and time.perf_counter() < end:
            time.sleep(0.01)
        time.sleep(0.25)

        poses = list(self._last_path_msg.poses)
        plan, smooth = self._alog["plan"], self._alog["smooth"]
        return {"success": len(poses) > 0, "path_length_m": path_length(poses),
                "num_waypoints": len(poses), "search_time_s": plan, "smoothing_time_s": smooth,
                "total_time_s": (plan + (smooth or 0.0)) if plan is not None else None,
                "expansions": self._alog["expansions"], "wall_time_s": wall, "poses": poses}

    def _publish_batch_goal(self, start, goal):
        self._reset_yours_state()
        self.broadcast_start_tf(*start)
        time.sleep(START_SETTLE_S)
        self._batch_publishing = True
        self._goal_t0 = time.perf_counter()
        self.goal_pub.publish(self.make_pose(*goal))

    # ------------------------------------------------------------------
    # One trial: plan, log, visualize
    # ------------------------------------------------------------------
    def run_comparison(self, label, start, goal, batch):
        self.trial_counter += 1
        n = self.trial_counter
        self.get_logger().info(f"--- Trial {n} ({label}) start={tuple(round(v, 2) for v in start)} "
                               f"goal={tuple(round(v, 2) for v in goal)}")

        # a_star_node ignores a goal identical to its last successful one (prev_goal guard)
        gkey = tuple(round(v, 9) for v in goal)
        if gkey == self._last_yours_goal_ok:
            r_yours = {"success": False, "error": "skipped: a_star_node ignores a goal identical to its last one"}
        else:
            if batch:
                self._publish_batch_goal(start, goal)
            r_yours = self.collect_yours()
            if r_yours["success"]:
                self._last_yours_goal_ok = gkey
        self._batch_publishing = False
        self.get_logger().info(f"  yours: { {k: v for k, v in r_yours.items() if k != 'poses'} }")

        r_nav2 = self.run_nav2(start, goal)      # after yours, so they never compete for CPU
        self.get_logger().info(f"  nav2 : { {k: v for k, v in r_nav2.items() if k != 'poses'} }")

        cid, snap = self.snapshot_config()

        fallback = r_yours.get("total_time_s") is None
        y_time = r_yours.get("wall_time_s") if fallback else r_yours["total_time_s"]
        n_time = r_nav2.get("planning_time_s")
        row = {
            "trial": n, "label": label, "timestamp": datetime.now().isoformat(timespec="seconds"),
            "config_id": cid,
            "start_x": rnd(start[0], 3), "start_y": rnd(start[1], 3), "start_yaw_deg": rnd(math.degrees(start[2]), 1),
            "goal_x": rnd(goal[0], 3), "goal_y": rnd(goal[1], 3), "goal_yaw_deg": rnd(math.degrees(goal[2]), 1),
            "straight_line_m": rnd(math.hypot(goal[0] - start[0], goal[1] - start[1]), 3),
            "nav2_success": bool(r_nav2.get("success")),
            "nav2_path_length_m": rnd(r_nav2.get("path_length_m"), 3),
            "nav2_num_waypoints": r_nav2.get("num_waypoints", ""),
            "nav2_planning_time_s": rnd(n_time, 4), "nav2_wall_time_s": rnd(r_nav2.get("wall_time_s"), 4),
            "nav2_error": r_nav2.get("error", ""),
            "yours_success": bool(r_yours.get("success")),
            "yours_path_length_m": rnd(r_yours.get("path_length_m"), 3),
            "yours_num_waypoints": r_yours.get("num_waypoints", ""),
            "yours_search_time_s": rnd(r_yours.get("search_time_s"), 4),
            "yours_smoothing_time_s": rnd(r_yours.get("smoothing_time_s"), 4),
            "yours_total_time_s": rnd(y_time, 4), "yours_wall_time_s": rnd(r_yours.get("wall_time_s"), 4),
            "yours_time_is_wallclock_fallback": fallback if r_yours.get("success") else "",
            "yours_expansions": r_yours.get("expansions", ""),
            "yours_error": r_yours.get("error", ""),
        }
        if r_nav2.get("path_length_m") and r_yours.get("path_length_m"):
            row["length_diff_m_yours_minus_nav2"] = round(r_yours["path_length_m"] - r_nav2["path_length_m"], 3)
        if n_time is not None and y_time is not None and r_yours.get("success"):
            row["planning_time_diff_s_yours_minus_nav2"] = round(y_time - n_time, 4)

        path_rows = []
        for planner, r in (("nav2", r_nav2), ("yours", r_yours)):
            for i, p in enumerate(r.get("poses", [])):
                path_rows.append([n, planner, i, round(p.pose.position.x, 4), round(p.pose.position.y, 4),
                                  round(math.degrees(quat_to_yaw(p.pose.orientation)), 2)])

        saved = self._save(row, path_rows, cid, snap)
        self._publish_result_markers(goal, r_nav2, r_yours, y_time)
        if saved:
            self.get_logger().info(f"  logged trial {n} ({len(path_rows)} path points, config {cid})")

    def run_batch(self):
        for name, sc in CUSTOM_RUNS.items():
            self._publish_start_marker(*sc["start"])
            self._publish_goal_marker(*sc["goal"])
            self.run_comparison(name, sc["start"], sc["goal"], batch=True)
            time.sleep(1.0)

    def _worker(self):
        if CUSTOM_RUNS:
            self._busy.set()
            try:
                self.get_logger().info(f"Running {len(CUSTOM_RUNS)} batch pair(s) first...")
                self.run_batch()
            except Exception:
                self.get_logger().error(traceback.format_exc())
            finally:
                self._busy.clear()
            self.get_logger().info("Batch done. Interactive mode.")
        while rclpy.ok() and not self._stop.is_set():
            try:
                job = self._jobs.get(timeout=0.5)
            except queue.Empty:
                continue
            try:
                self.run_comparison(*job)
            except Exception:
                self.get_logger().error(traceback.format_exc())
            finally:
                self._busy.clear()

    # ------------------------------------------------------------------
    # RViz markers
    # ------------------------------------------------------------------
    def _clear_markers(self):
        arr = MarkerArray()
        m = Marker()
        m.header.frame_id = MAP_FRAME
        m.action = Marker.DELETEALL
        arr.markers.append(m)
        self.marker_pub.publish(arr)

    def _arrow(self, ns, x, y, yaw, color):
        m = Marker()
        m.header.frame_id = MAP_FRAME
        m.header.stamp = self.get_clock().now().to_msg()
        m.ns, m.id = ns, self.trial_counter + 1
        m.type, m.action = Marker.ARROW, Marker.ADD
        m.pose.position.x, m.pose.position.y, m.pose.position.z = x, y, 0.05
        qx, qy, qz, qw = yaw_to_quat(yaw)
        m.pose.orientation = Quaternion(x=qx, y=qy, z=qz, w=qw)
        m.scale.x, m.scale.y, m.scale.z = 0.5, 0.12, 0.12
        m.color = color
        m.lifetime = Duration(seconds=0).to_msg()
        return m

    def _publish_start_marker(self, x, y, yaw):
        self.marker_pub.publish(MarkerArray(markers=[self._arrow("compare_start", x, y, yaw, rgba(0.0, 0.85, 0.0, 0.9))]))

    def _publish_goal_marker(self, x, y, yaw):
        self.marker_pub.publish(MarkerArray(markers=[self._arrow("compare_goal", x, y, yaw, rgba(0.9, 0.0, 0.0, 0.9))]))

    def _path_marker(self, ns, poses, color):
        m = Marker()
        m.header.frame_id = MAP_FRAME
        m.header.stamp = self.get_clock().now().to_msg()
        m.ns, m.id = ns, 0
        m.type, m.action = Marker.LINE_STRIP, Marker.ADD
        m.scale.x = 0.06
        m.color = color
        m.lifetime = Duration(seconds=0).to_msg()
        m.points = [Point(x=p.pose.position.x, y=p.pose.position.y, z=0.03) for p in poses]
        return m

    def _publish_result_markers(self, goal, r_nav2, r_yours, y_time):
        arr = MarkerArray()
        # always publish both path markers; an empty one clears the previous trial's line
        arr.markers.append(self._path_marker("nav2_path", r_nav2.get("poses", []), rgba(1.0, 0.55, 0.0, 0.95)))
        arr.markers.append(self._path_marker("yours_path", r_yours.get("poses", []), rgba(0.0, 0.6, 1.0, 0.95)))
        nav2_line = (f"Nav2: {fmt(r_nav2.get('path_length_m'), 2)} m / {fmt(r_nav2.get('planning_time_s'))} s"
                     if r_nav2.get("success") else f"Nav2: FAILED ({r_nav2.get('error')})")
        yours_line = (f"Yours: {fmt(r_yours.get('path_length_m'), 2)} m / {fmt(y_time)} s"
                      if r_yours.get("success") else f"Yours: FAILED ({r_yours.get('error')})")
        t = Marker()
        t.header.frame_id = MAP_FRAME
        t.header.stamp = self.get_clock().now().to_msg()
        t.ns, t.id = "result_text", 0
        t.type, t.action = Marker.TEXT_VIEW_FACING, Marker.ADD
        t.pose.position.x, t.pose.position.y, t.pose.position.z = goal[0], goal[1], 0.8
        t.scale.z = 0.35
        t.color = rgba(1.0, 1.0, 1.0, 1.0)
        t.lifetime = Duration(seconds=0).to_msg()
        t.text = f"[orange] {nav2_line}\n[blue] {yours_line}"
        arr.markers.append(t)
        self.marker_pub.publish(arr)


def main():
    rclpy.init()
    node = CompareNode()
    executor = MultiThreadedExecutor(num_threads=4)
    executor.add_node(node)
    node.start()
    try:
        executor.spin()
    except KeyboardInterrupt:
        pass
    finally:
        node.get_logger().info(f"Logged up to trial {node.trial_counter} in {EXCEL_PATH}")
        node.stop()
        executor.shutdown()
        node.destroy_node()
        rclpy.shutdown()


if __name__ == "__main__":
    main()